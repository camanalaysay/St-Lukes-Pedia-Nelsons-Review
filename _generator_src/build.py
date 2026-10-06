#!/usr/bin/env python3
"""Build the self-contained St. Luke's Pediatrics Exam Generator HTML.

Reads the Nelson's question bank workbook, embeds the questions, references,
St. Luke's logo and jsPDF into template.html, and writes one offline HTML file.

Usage:  python3 build.py [path/to/QuestionBank.xlsx]
"""
import base64, collections, glob, json, os, re, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, "20261006_PED_NelsonsExamGenerator_v1.0.html")

xlsx = sys.argv[1] if len(sys.argv) > 1 else sorted(glob.glob(os.path.join(ROOT, "*NelsonsQuestionBank*.xlsx")))[-1]
wb = openpyxl.load_workbook(xlsx, read_only=True)

DIFF = ["Easy", "Average", "Difficult"]
COG = ["Recall", "Application", "Clinical Reasoning"]
CROSS_REF = re.compile(r"\b(all|none|both) of the (above|these|options)\b|\b[A-D] and [A-D]\b|\b(options?|choices?) [A-D]\b", re.I)

BOLD_MD = re.compile(r"\*\*([^*]+)\*\*")  # source stems emphasize words as **MOST**; they are already capitalized

def s(v):
    return "" if v is None else BOLD_MD.sub(r"\1", str(v)).strip()

# Full citations, deduplicated
cites, cite_idx, q_refs = [], {}, collections.defaultdict(list)
for r in wb["References"].iter_rows(min_row=5, values_only=True):
    if not r[0]:
        continue
    c = s(r[5])
    if c not in cite_idx:
        cite_idx[c] = len(cites)
        cites.append(c)
    q_refs[r[0]].append((r[2] or 0, cite_idx[c]))

sections = {}
for r in wb["Sections"].iter_rows(min_row=5, values_only=True):
    if isinstance(r[0], int):
        sections[r[0]] = {"no": r[0], "title": s(r[1]), "minutes": r[4], "n": 0}

qs, stripped_marks, skipped = [], 0, collections.Counter()
for r in wb["Question Bank"].iter_rows(min_row=5, values_only=True):
    if not r[0]:
        continue
    status = s(r[23])
    if status.lower() == "retired":
        skipped["Retired"] += 1
        continue
    opts = []
    for o in r[4:8]:
        o = s(o)
        if "✓" in o:  # answer tick left in some source options would give the key away
            stripped_marks += 1
            o = o.replace("✓", "").strip()
        opts.append(o)
    ans = s(r[8]).upper()
    if ans not in "ABCD" or not ans or not s(r[3]):
        skipped["Incomplete"] += 1
        continue
    sec = int(r[1])
    sections[sec]["n"] += 1
    refs = [i for _, i in sorted(q_refs.get(r[0], []))]
    qs.append([
        s(r[0]), sec, int(r[2] or 0), s(r[3]), opts, "ABCD".index(ans), s(r[9]),
        [s(x) for x in r[10:14]], s(r[14]),
        DIFF.index(s(r[15])) if s(r[15]) in DIFF else 1,
        COG.index(s(r[16])) if s(r[16]) in COG else 0,
        1 if "K" in s(r[17]) else 0,
        1 if any(CROSS_REF.search(o) for o in opts) else 0,
        status, refs,
    ])

data = {
    "source": os.path.basename(xlsx),
    "sections": [sections[k] for k in sorted(sections)],
    "q": qs,
    "cites": cites,
}

logo = base64.b64encode(open(os.path.join(HERE, "stlukes-logo.png"), "rb").read()).decode()
html = open(os.path.join(HERE, "template.html"), encoding="utf-8").read()
html = (html.replace("/*__JSPDF__*/", open(os.path.join(HERE, "jspdf.min.js"), encoding="utf-8").read())
            .replace("__LOGO__", "data:image/png;base64," + logo)
            .replace("/*__DATA__*/null", json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")))
open(OUT, "w", encoding="utf-8").write(html)

print(f"{len(qs)} questions, {len(cites)} citations, {stripped_marks} answer ticks removed, skipped {dict(skipped)}")
print(f"wrote {OUT} ({os.path.getsize(OUT)/1e6:.1f} MB)")
